"""Regenerate the hand-designed ingestion fixtures.

Each fixture isolates one failure mode of naive spreadsheet reading. The files are
committed; this script exists so the shapes stay readable and reproducible.

    python backend/tests/fixtures/make_fixtures.py
"""

from __future__ import annotations

import pathlib

from openpyxl import Workbook

HERE = pathlib.Path(__file__).parent


def write_grid(ws, grid: list[list[object]]) -> None:
    for row in grid:
        ws.append(row)


def clean_csv() -> None:
    (HERE / "clean.csv").write_text(
        "order_id,order_date,amount\n"
        "ORD-1001,2024-01-05,120.50\n"
        "ORD-1002,2024-01-07,98.00\n"
        "ORD-1003,2024-02-11,410.25\n"
        "ORD-1004,2024-02-19,65.75\n"
        "ORD-1005,2024-03-02,1200.00\n"
    )


def no_header_csv() -> None:
    # Every row is data: the first row must NOT be mistaken for a header.
    (HERE / "no_header.csv").write_text(
        "C-001,2024-01-05,120.50\n"
        "C-002,2024-01-07,98.00\n"
        "C-003,2024-02-11,410.25\n"
        "C-004,2024-02-19,65.75\n"
        "C-005,2024-03-02,1200.00\n"
    )


def two_tables_one_sheet() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    write_grid(
        ws,
        [
            ["order_id", "region", "amount"],
            ["ORD-1001", "North", 120.5],
            ["ORD-1002", "South", 98.0],
            ["ORD-1003", "North", 410.25],
            ["ORD-1004", "East", 65.75],
            [None, None, None],
            ["emp_id", "dept"],
            ["EMP-01", "Sales"],
            ["EMP-02", "Ops"],
            ["EMP-03", "Sales"],
        ],
    )
    wb.save(HERE / "two_tables_one_sheet.xlsx")


def title_and_footnote() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Q3"
    write_grid(
        ws,
        [
            ["Quarterly Revenue Report", None, None],
            ["Generated 2024-10-01", None, None],
            [None, None, None],
            ["region", "quarter", "revenue"],
            ["North", "Q3", 12000.0],
            ["South", "Q3", 9800.5],
            ["East", "Q3", 15410.0],
            ["West", "Q3", 11000.0],
            ["Total", None, 48210.5],
            ["Source: internal ledger", None, None],
        ],
    )
    wb.save(HERE / "title_and_footnote.xlsx")


def merged_header() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll"
    write_grid(
        ws,
        [
            ["Employee Details", None, "Salary"],
            ["Asha Menon", "Sales", 82000.0],
            ["Ravi Kumar", "Ops", 74000.0],
            ["Lin Wei", "Sales", 91000.0],
            ["Diego Alvarez", "Finance", 68000.0],
        ],
    )
    ws.merge_cells("A1:B1")
    wb.save(HERE / "merged_header.xlsx")


if __name__ == "__main__":
    clean_csv()
    no_header_csv()
    two_tables_one_sheet()
    title_and_footnote()
    merged_header()
    print(f"wrote fixtures to {HERE}")
